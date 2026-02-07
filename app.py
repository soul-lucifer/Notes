from flask import Flask, request, jsonify, render_template_string
import os
import PyPDF2
from pptx import Presentation
import google.generativeai as genai
from werkzeug.utils import secure_filename

app = Flask(__name__)

# Gemini API setup
genai.configure(api_key=os.environ.get('GOOGLE_API_KEY'))

model = genai.GenerativeModel('gemini-1.5-flash')  # Free tier model

ALLOWED_EXTENSIONS = {'pdf', 'pptx'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_pdf(file_path):
    with open(file_path, 'rb') as file:
        reader = PyPDF2.PdfReader(file)
        text = ''
        for page in reader.pages:
            text += page.extract_text() + '\n'
        return text

def extract_text_from_pptx(file_path):
    prs = Presentation(file_path)
    text = ''
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, 'text'):
                text += shape.text + '\n'
    return text

def generate_notes(text):
    prompt = (
        "You are a helpful assistant that generates systematic notes from document content. "
        "Structure the notes with headings, bullet points, key insights, and summaries.\n\n"
        f"Generate systematic notes from this content: {text[:8000]}"  # Limit to avoid token limits; expand as needed
    )
    response = model.generate_content(prompt)
    return response.text

@app.route('/')
def index():
    return render_template_string(open('index.html').read())

@app.route('/generate', methods=['POST'])
def generate():
    if 'file' not in request.files:
        return jsonify({'error': 'No file part'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No selected file'}), 400
    if file and allowed_file(file.filename):
        filename = secure_filename(file.filename)
        file_path = os.path.join('/tmp', filename)  # Vercel uses /tmp for temp files
        file.save(file_path)
        
        try:
            if filename.endswith('.pdf'):
                text = extract_text_from_pdf(file_path)
            elif filename.endswith('.pptx'):
                text = extract_text_from_pptx(file_path)
            else:
                return jsonify({'error': 'Unsupported file type'}), 400
            
            notes = generate_notes(text)
            os.remove(file_path)
            return jsonify({'notes': notes})
        except Exception as e:
            os.remove(file_path)
            return jsonify({'error': str(e)}), 500
    return jsonify({'error': 'Invalid file'}), 400

if __name__ == '__main__':
    app.run()
